"""Pruebas de app/export.py (F6 / T2): antes no habia ninguna."""
import os
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.export import (
    build_epub_bytes,
    build_markdown,
    build_txt,
    chapter_heading,
    chapter_paragraphs,
    clean_chapter_text,
    export_docx,
    export_epub,
    export_markdown,
    export_txt,
    parse_inline,
)
from app.project import Project


def _project():
    p = Project(name="Mi novela", author="Ana Autora")
    ch = p.add_chapter("The Storm", 100)
    ch.content = "First *thought here* and **loud** words.\n\nSecond paragraph."
    p.add_chapter("Aftermath", 100).content = "Quiet."
    return p


class TestInlineMarkup(unittest.TestCase):
    def test_italic(self):
        self.assertEqual(parse_inline("a *b c* d"), [("a ", False, False), ("b c", False, True), (" d", False, False)])

    def test_bold_is_not_turned_into_italic(self):
        segs = parse_inline("**Bold** text")
        self.assertEqual(segs[0], ("Bold", True, False))
        self.assertNotIn("*", "".join(s[0] for s in segs))

    def test_spaced_asterisks_are_not_italic(self):
        segs = parse_inline("a * b and c * d")
        self.assertEqual(segs, [("a * b and c * d", False, False)])

    def test_unbalanced_asterisk_is_kept_literally(self):
        self.assertEqual(parse_inline("5 *stars"), [("5 *stars", False, False)])

    def test_italic_inside_bold(self):
        segs = parse_inline("**very *deep* bold**")
        self.assertIn(("deep", True, True), segs)


class TestParagraphs(unittest.TestCase):
    def test_italic_spanning_two_paragraphs_is_closed_in_each(self):
        paras = chapter_paragraphs("*I can't.\n\nI won't.* She left.")
        self.assertEqual(paras, ["*I can't.*", "*I won't.* She left."])
        self.assertTrue(parse_inline(paras[0])[0][2])
        self.assertTrue(parse_inline(paras[1])[0][2])

    def test_model_emitted_headings_are_removed(self):
        text = "## Chapter 3: The Storm\n\nIt rained.\n\n### A Break\n\nMore."
        self.assertEqual(clean_chapter_text(text), "It rained.\n\nA Break\n\nMore.")

    def test_heading_label_is_configurable(self):
        p = _project()
        self.assertEqual(chapter_heading(p, 1, p.chapters[0]), "Chapter 1: The Storm")
        p.chapter_heading = "Capítulo"
        self.assertEqual(chapter_heading(p, 1, p.chapters[0]), "Capítulo 1: The Storm")
        p.chapter_heading = ""
        self.assertEqual(chapter_heading(p, 1, p.chapters[0]), "The Storm")


class TestTextExports(unittest.TestCase):
    def test_txt_has_title_author_and_chapters(self):
        txt = build_txt(_project())
        self.assertTrue(txt.startswith("Mi novela\n========="))
        self.assertIn("Ana Autora", txt)
        self.assertIn("Chapter 2: Aftermath", txt)

    def test_markdown(self):
        md = build_markdown(_project())
        self.assertIn("# Mi novela", md)
        self.assertIn("## Chapter 1: The Storm", md)
        self.assertIn("*thought here*", md)

    def test_writes_are_atomic_and_leave_no_tmp(self):
        with tempfile.TemporaryDirectory() as tmp:
            for fn, name in ((export_txt, "a.txt"), (export_markdown, "a.md")):
                path = os.path.join(tmp, name)
                fn(_project(), path)
                self.assertTrue(os.path.getsize(path) > 0)
                self.assertFalse(os.path.exists(path + ".tmp"))


class TestDocx(unittest.TestCase):
    def test_runs_page_breaks_title_and_properties(self):
        from docx import Document
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "story.docx")
            export_docx(_project(), path)
            doc = Document(path)
        self.assertEqual(doc.core_properties.title, "Mi novela")
        self.assertEqual(doc.core_properties.author, "Ana Autora")
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("Chapter 1: The Storm", texts)
        body = next(p for p in doc.paragraphs if p.text.startswith("First"))
        self.assertEqual(body.text, "First thought here and loud words.")
        self.assertTrue(any(r.italic and r.text == "thought here" for r in body.runs))
        self.assertTrue(any(r.bold and r.text == "loud" for r in body.runs))
        breaks = [br for p in doc.paragraphs for br in p._p.iter(qn("w:br")) if br.get(qn("w:type")) == "page"]
        self.assertEqual(len(breaks), 2)  # uno antes de cada capitulo


class TestEpub(unittest.TestCase):
    def test_epub_structure(self):
        data = build_epub_bytes(_project())
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "story.epub")
            export_epub(_project(), path)
            with zipfile.ZipFile(path) as zf:
                names = zf.namelist()
                self.assertEqual(names[0], "mimetype")
                self.assertEqual(zf.getinfo("mimetype").compress_type, zipfile.ZIP_STORED)
                self.assertEqual(zf.read("mimetype"), b"application/epub+zip")
                self.assertIn("OEBPS/chapter_01.xhtml", names)
                ch1 = zf.read("OEBPS/chapter_01.xhtml").decode("utf-8")
                self.assertIn("<em>thought here</em>", ch1)
                self.assertIn("<strong>loud</strong>", ch1)
                opf = zf.read("OEBPS/content.opf").decode("utf-8")
                self.assertIn("<dc:creator>Ana Autora</dc:creator>", opf)
        self.assertTrue(data.startswith(b"PK"))

    def test_xhtml_is_well_formed_and_escaped(self):
        import xml.dom.minidom

        p = Project(name="A & B <x>")
        p.add_chapter("T", 10).content = "Tom & Jerry said \"<hi>\"."
        with zipfile.ZipFile(__import__("io").BytesIO(build_epub_bytes(p))) as zf:
            for name in zf.namelist():
                if name.endswith((".xhtml", ".opf", ".ncx", ".xml")):
                    xml.dom.minidom.parseString(zf.read(name))  # no lanza


if __name__ == "__main__":
    unittest.main()
