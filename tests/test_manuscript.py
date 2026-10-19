"""Importar manuscrito (F16): .txt/.md/.docx split into chapters by their headings, the chapter
limit, the desktop menu action and the web API."""
import base64
import http.client
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from test_web_server import FakeEngine  # noqa: E402

from app.manuscript import read_manuscript  # noqa: E402
from app.project import MAX_CHAPTERS, Project  # noqa: E402
from app.web_server import ApiError, StoryForgeService, create_server  # noqa: E402

PARAGRAPH = "The lighthouse keeper climbed the stairs again, counting each step as the storm rose. "


def _docx(blocks):
    """blocks: ("h1"|"h2"|"title", text) or ("p", [(text, bold, italic), ...])."""
    from docx import Document

    document = Document()
    for kind, value in blocks:
        if kind == "title":
            document.add_heading(value, level=0)
        elif kind in ("h1", "h2"):
            document.add_heading(value, level=int(kind[1]))
        else:
            paragraph = document.add_paragraph()
            for text, bold, italic in value:
                run = paragraph.add_run(text)
                run.bold, run.italic = bold, italic
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestPlainText(unittest.TestCase):
    def test_chapter_headings_in_english_and_spanish(self):
        text = "\n".join([
            "The Keeper", "", "Chapter 1: The Storm", "", PARAGRAPH, "", PARAGRAPH,
            "CAPÍTULO II", PARAGRAPH, "Capítulo tres — El faro", PARAGRAPH, "Epilogue", "The end came quietly.",
        ])
        manuscript = read_manuscript(text.encode("utf-8"), "keeper.txt")
        self.assertEqual(manuscript.title, "The Keeper")  # a short line before the first chapter
        self.assertEqual([s.title for s in manuscript.sections],
                         ["Chapter 1: The Storm", "CAPÍTULO II", "Capítulo tres — El faro", "Epilogue"])
        project = manuscript.to_project()
        self.assertEqual(project.chapters[0].content, f"{PARAGRAPH.strip()}\n\n{PARAGRAPH.strip()}")
        self.assertEqual(project.chapters[3].content, "The end came quietly.")
        self.assertEqual(project.chapters[0].target_words, 2000)

    def test_long_text_before_the_first_heading_becomes_a_chapter(self):
        text = (PARAGRAPH * 5) + "\nChapter 1\n" + PARAGRAPH + "\nChapter 2\n" + PARAGRAPH
        manuscript = read_manuscript(text.encode("cp1252"), "novel.txt")
        self.assertEqual([s.title for s in manuscript.sections], ["Inicio", "Chapter 1", "Chapter 2"])
        self.assertEqual(manuscript.title, "novel")

    def test_no_headings_gives_one_chapter_and_empty_is_refused(self):
        manuscript = read_manuscript((PARAGRAPH * 3).encode(), "draft.txt")
        self.assertEqual(len(manuscript.sections), 1)
        with self.assertRaises(ValueError):
            read_manuscript(b"   \n\n ", "empty.txt")
        with self.assertRaises(ValueError):
            read_manuscript(b"x", "book.pdf")

    def test_more_chapters_than_the_limit_are_merged_into_the_last(self):
        text = "\n".join(f"Chapter {n}\nText of chapter {n}." for n in range(1, MAX_CHAPTERS + 4))
        manuscript = read_manuscript(text.encode(), "long.txt")
        self.assertEqual(len(manuscript.sections), MAX_CHAPTERS)
        self.assertEqual(manuscript.merged, 3)
        last = manuscript.to_project().chapters[-1].content
        self.assertIn(f"Text of chapter {MAX_CHAPTERS}.", last)
        self.assertIn(f"Chapter {MAX_CHAPTERS + 3}\n\nText of chapter {MAX_CHAPTERS + 3}.", last)  # nothing lost


class TestMarkdownAndDocx(unittest.TestCase):
    def test_markdown_book_title_above_chapters(self):
        text = "# My Novel\n\n## One\n\nFirst *chapter*.\n\n### A scene break\n\nMore.\n\n## Two\n\nSecond.\n"
        manuscript = read_manuscript(text.encode(), "novel.md")
        self.assertEqual(manuscript.title, "My Novel")
        self.assertEqual([s.title for s in manuscript.sections], ["One", "Two"])
        self.assertIn("A scene break", manuscript.sections[0].text)
        self.assertIn("First *chapter*.", manuscript.sections[0].text)

    def test_docx_headings_and_formatting(self):
        content = _docx([
            ("title", "Salt and Light"),
            ("h1", "Chapter One"),
            ("p", [("She ", False, False), ("never", False, True), (" looked ", False, False), ("back", True, False), (".", False, False)]),
            ("p", [("Split ", False, True), ("italic", False, True)]),
            ("h1", "Chapter Two"),
            ("p", [(PARAGRAPH, False, False)]),
        ])
        manuscript = read_manuscript(content, "salt.docx")
        self.assertEqual(manuscript.title, "Salt and Light")
        self.assertEqual([s.title for s in manuscript.sections], ["Chapter One", "Chapter Two"])
        self.assertEqual(manuscript.sections[0].lines[:2], ["She *never* looked **back**.", "*Split italic*"])

    def test_docx_without_heading_styles_falls_back_to_heading_lines(self):
        content = _docx([("p", [("Prologue", False, False)]), ("p", [("Before.", False, False)]),
                         ("p", [("Chapter 1", False, False)]), ("p", [("After.", False, False)])])
        self.assertEqual([s.title for s in read_manuscript(content, "plain.docx").sections], ["Prologue", "Chapter 1"])

    def test_a_broken_docx_is_a_clear_error(self):
        with self.assertRaises(ValueError):
            read_manuscript(b"PK\x03\x04 not really a docx", "broken.docx")


class TestWebImport(unittest.TestCase):
    def test_service_creates_and_opens_the_project(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        service = StoryForgeService(FakeEngine(), tmp.name)
        text = "Chapter 1\nOne.\nChapter 2\nTwo."
        view = service.import_manuscript("story.txt", base64.b64encode(text.encode()).decode())
        self.assertEqual((view["nombre"], len(view["capitulos"])), ("story", 2))
        self.assertIn("2 capítulos", view["mensaje"])
        saved = Project.load(os.path.join(tmp.name, view["archivo"]))
        self.assertEqual([c.title for c in saved.chapters], ["Chapter 1", "Chapter 2"])
        second = service.import_manuscript("story.txt", base64.b64encode(text.encode()).decode())
        self.assertNotEqual(second["archivo"], view["archivo"])  # never overwrites
        with self.assertRaises(ApiError):
            service.import_manuscript("story.txt", "%%%not base64")
        with self.assertRaises(ApiError):
            service.import_manuscript("story.pdf", base64.b64encode(b"x").decode())

    def test_http_route_accepts_a_body_over_the_normal_limit(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        server = create_server(FakeEngine(), tmp.name, port=0, host="127.0.0.1", token="k" * 24)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        conn = http.client.HTTPConnection("127.0.0.1", server.server_address[1], timeout=10)
        self.addCleanup(conn.close)
        big = ("Chapter 1\n" + PARAGRAPH * 25000 + "\nChapter 2\nEnd.").encode()  # ~2.2 MB of text
        body = json.dumps({"nombre": "big.txt", "datos": base64.b64encode(big).decode()})
        conn.request("POST", "/api/proyectos/importar", body=body,
                     headers={"Content-Type": "application/json", "X-StoryForge-Token": "k" * 24})
        response = conn.getresponse()
        payload = json.loads(response.read())
        self.assertEqual(response.status, 200, payload)
        self.assertEqual(len(payload["capitulos"]), 2)


class TestDesktopMenu(unittest.TestCase):
    def test_import_sets_an_unsaved_project(self):
        from app import gui

        app = mock.MagicMock()
        app.busy = False
        app._confirm_discard_changes.return_value = True
        folder = tempfile.mkdtemp()
        path = os.path.join(folder, "book.md")
        with open(path, "w", encoding="utf-8") as handle:
            handle.write("## One\nFirst.\n## Two\nSecond.\n")
        with mock.patch.object(gui.messagebox, "showinfo") as info:
            self.assertTrue(gui.StoryForgeApp.import_manuscript(app, path))
        self.assertEqual([c.title for c in app.project.chapters], ["One", "Two"])
        self.assertIsNone(app.current_path)
        app._load_into_ui.assert_called_once()
        self.assertIn("2 capítulos", info.call_args[0][1])


if __name__ == "__main__":
    unittest.main()
