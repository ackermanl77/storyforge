"""
Exportacion de la historia a .docx, .txt, Markdown (.md) y EPUB.

- Pensamientos internos (*texto*) se exportan como cursiva real y **texto**
  como negrita (en ese orden: antes "**Bold**" salia como "*Bold" en
  cursiva con un asterisco suelto). Un asterisco solo cuenta como marca si
  no tiene un espacio del lado de adentro ("a * b and c * d" no es cursiva),
  y una cursiva que abarca dos parrafos se cierra y se reabre en cada uno.
- Los encabezados markdown que el modelo a veces escribe ("# Chapter 3",
  "## La tormenta") se limpian: el capitulo ya tiene su propio encabezado.
- .docx: portada (titulo y autor), un salto de pagina antes de cada capitulo
  y propiedades del documento (titulo, autor).
- EPUB 3 hecho a mano (zip + XHTML), sin dependencias extra; incluye toc.ncx
  para lectores EPUB 2.
- El encabezado de cada capitulo es configurable (Project.chapter_heading:
  "Chapter", "Capítulo" o "" para solo el titulo).
- Todas las escrituras son atomicas (app/fileutil.py).
"""
from __future__ import annotations

import html
import io
import re
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Tuple

from .fileutil import atomic_write_bytes, atomic_write_text, atomic_write_with
from .project import Chapter, Project

# Un segmento de texto: (texto, negrita, cursiva).
Segment = Tuple[str, bool, bool]

_BOLD_RE = re.compile(r"(?<!\*)\*\*(?=\S)(.+?)(?<=\S)\*\*(?!\*)")
_ITALIC_RE = re.compile(r"(?<![*\w])\*(?=[^\s*])(.+?)(?<=[^\s*])\*(?![*\w])")
# Cursiva que empieza en un parrafo y termina en otro (como mucho 3 parrafos).
_CROSS_PARAGRAPH_ITALIC_RE = re.compile(
    r"(?<![*\w])\*(?=[^\s*])((?:(?!\n[ \t]*\n).)+?(?:\n[ \t]*\n(?:(?!\n[ \t]*\n).)+?){1,2})(?<=[^\s*])\*(?![*\w])",
    re.DOTALL,
)
_MD_HEADING_RE = re.compile(r"^[ \t]*#{1,6}[ \t]+(.*?)[ \t#]*$")
_CHAPTER_HEADING_LINE_RE = re.compile(
    r"^[ \t]*(?:#{1,6}[ \t]*)?(?:\*\*|__)?[ \t]*(?:chapter|cap[ií]tulo)\b[^\n]{0,80}$", re.IGNORECASE,
)
_PARAGRAPH_SPLIT_RE = re.compile(r"\n[ \t]*\n")


# ---------------------------------------------------------------- limpieza
def clean_chapter_text(text: str) -> str:
    """Quita encabezados markdown que el modelo haya escrito: una primera
    linea tipo "Chapter 3: ..." se elimina (el exportador pone el suyo) y en
    las demas se quitan los "#"."""
    lines = (text or "").strip().splitlines()
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and _CHAPTER_HEADING_LINE_RE.match(lines[0]):
        lines.pop(0)
    cleaned = []
    for line in lines:
        match = _MD_HEADING_RE.match(line)
        cleaned.append(match.group(1) if match else line)
    return "\n".join(cleaned).strip()


def _close_italics_per_paragraph(text: str) -> str:
    def fix(match: re.Match) -> str:
        inner = _PARAGRAPH_SPLIT_RE.sub(lambda m: "*" + m.group(0) + "*", match.group(1))
        return f"*{inner}*"
    return _CROSS_PARAGRAPH_ITALIC_RE.sub(fix, text)


def chapter_paragraphs(text: str) -> List[str]:
    """Parrafos (separados por linea en blanco) ya limpios y con las cursivas
    que cruzaban parrafos cerradas en cada uno."""
    text = _close_italics_per_paragraph(clean_chapter_text(text))
    paragraphs = []
    for para in _PARAGRAPH_SPLIT_RE.split(text):
        para = " ".join(line.strip() for line in para.strip().splitlines())
        if para:
            paragraphs.append(para)
    return paragraphs


def parse_inline(text: str) -> List[Segment]:
    """Parte un parrafo en segmentos (texto, negrita, cursiva)."""
    segments: List[Segment] = []

    def italics(chunk: str, bold: bool) -> None:
        pos = 0
        for m in _ITALIC_RE.finditer(chunk):
            if m.start() > pos:
                segments.append((chunk[pos:m.start()], bold, False))
            segments.append((m.group(1), bold, True))
            pos = m.end()
        if pos < len(chunk):
            segments.append((chunk[pos:], bold, False))

    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            italics(text[pos:m.start()], False)
        italics(m.group(1), True)
        pos = m.end()
    if pos < len(text):
        italics(text[pos:], False)
    return [s for s in segments if s[0]]


def chapter_heading(project: Project, index: int, chapter: Chapter) -> str:
    label = (project.chapter_heading or "").strip()
    title = chapter.title.strip()
    if not label:
        return title or str(index)
    return f"{label} {index}: {title}" if title else f"{label} {index}"


def _title(project: Project) -> str:
    return project.name.strip() or "Historia"


# ---------------------------------------------------------------- .txt / .md
def build_txt(project: Project) -> str:
    title = _title(project)
    lines = [title, "=" * len(title)]
    if project.author.strip():
        lines.append(project.author.strip())
    lines.append("")
    for i, chapter in enumerate(project.chapters, start=1):
        heading = chapter_heading(project, i, chapter)
        lines.append(heading)
        lines.append("-" * max(len(heading), 10))
        lines.append("")
        lines.append("\n\n".join(chapter_paragraphs(chapter.content)))
        lines.append("")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def export_txt(project: Project, path: str | Path) -> None:
    atomic_write_text(path, build_txt(project))


def build_markdown(project: Project) -> str:
    parts = [f"# {_title(project)}"]
    if project.author.strip():
        parts.append(f"*{project.author.strip()}*")
    for i, chapter in enumerate(project.chapters, start=1):
        parts.append(f"## {chapter_heading(project, i, chapter)}")
        parts.extend(chapter_paragraphs(chapter.content))
    return "\n\n".join(parts) + "\n"


def export_markdown(project: Project, path: str | Path) -> None:
    atomic_write_text(path, build_markdown(project))


# ---------------------------------------------------------------- .docx
def export_docx(project: Project, path: str | Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    except ImportError as exc:
        raise RuntimeError(
            "El paquete 'python-docx' no está instalado. "
            "Instala las dependencias con: pip install -r requirements.txt"
        ) from exc

    doc = Document()
    props = doc.core_properties
    props.title = _title(project)
    if project.author.strip():
        props.author = project.author.strip()

    # Portada
    title = doc.add_heading(_title(project), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if project.author.strip():
        author = doc.add_paragraph(project.author.strip())
        author.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, chapter in enumerate(project.chapters, start=1):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(chapter_heading(project, i, chapter), level=1)
        for paragraph_text in chapter_paragraphs(chapter.content):
            p = doc.add_paragraph()
            _add_runs(p, paragraph_text)

    atomic_write_with(path, lambda tmp: doc.save(tmp))


def _add_runs(paragraph, text: str) -> None:
    for chunk, bold, italic in parse_inline(text):
        run = paragraph.add_run(chunk)
        if bold:
            run.bold = True
        if italic:
            run.italic = True


# Compatibilidad con el nombre anterior.
_add_runs_with_italics = _add_runs


# ---------------------------------------------------------------- EPUB
def _inline_html(text: str) -> str:
    out = []
    for chunk, bold, italic in parse_inline(text):
        piece = html.escape(chunk, quote=False)
        if italic:
            piece = f"<em>{piece}</em>"
        if bold:
            piece = f"<strong>{piece}</strong>"
        out.append(piece)
    return "".join(out)


def _xhtml(title: str, body: str, lang: str) -> str:
    return (
        '<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
        f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
        f'lang="{lang}" xml:lang="{lang}">\n<head><meta charset="utf-8"/>'
        f"<title>{html.escape(title)}</title>"
        '<link rel="stylesheet" type="text/css" href="style.css"/></head>\n'
        f"<body>\n{body}\n</body>\n</html>\n"
    )


EPUB_CSS = """body { font-family: serif; line-height: 1.5; margin: 0 5%; }
h1, h2 { text-align: center; margin: 2em 0 1em; }
p { text-indent: 1.5em; margin: 0 0 0.3em; text-align: justify; }
p.author { text-align: center; text-indent: 0; font-style: italic; }
"""


def build_epub_bytes(project: Project, lang: str = "en") -> bytes:
    title = _title(project)
    author = project.author.strip() or "Anónimo"
    book_id = f"urn:uuid:{uuid.uuid4()}"
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    chapters = []
    for i, chapter in enumerate(project.chapters, start=1):
        heading = chapter_heading(project, i, chapter)
        paras = "\n".join(f"<p>{_inline_html(p)}</p>" for p in chapter_paragraphs(chapter.content))
        chapters.append((f"chapter_{i:02d}.xhtml", heading, _xhtml(heading, f"<h2>{html.escape(heading)}</h2>\n{paras}", lang)))

    title_body = f"<h1>{html.escape(title)}</h1>"
    if project.author.strip():
        title_body += f'\n<p class="author">{html.escape(project.author.strip())}</p>'
    title_page = _xhtml(title, title_body, lang)

    nav_items = "\n".join(
        f'<li><a href="{name}">{html.escape(heading)}</a></li>' for name, heading, _ in chapters
    )
    nav = _xhtml(
        title,
        f'<nav epub:type="toc" id="toc"><h1>{html.escape(title)}</h1><ol>\n'
        f'<li><a href="title.xhtml">{html.escape(title)}</a></li>\n{nav_items}\n</ol></nav>',
        lang,
    )
    ncx_points = "\n".join(
        f'<navPoint id="np{n}" playOrder="{n}"><navLabel><text>{html.escape(heading)}</text></navLabel>'
        f'<content src="{name}"/></navPoint>'
        for n, (name, heading, _) in enumerate(chapters, start=1)
    )
    ncx = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">\n'
        f'<head><meta name="dtb:uid" content="{book_id}"/></head>\n'
        f"<docTitle><text>{html.escape(title)}</text></docTitle>\n<navMap>\n{ncx_points}\n</navMap>\n</ncx>\n"
    )
    manifest = "\n".join(
        f'<item id="ch{n}" href="{name}" media-type="application/xhtml+xml"/>'
        for n, (name, _, _) in enumerate(chapters, start=1)
    )
    spine = "\n".join(f'<itemref idref="ch{n}"/>' for n in range(1, len(chapters) + 1))
    opf = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">\n'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        f'<dc:identifier id="bookid">{book_id}</dc:identifier>\n'
        f"<dc:title>{html.escape(title)}</dc:title>\n"
        f"<dc:creator>{html.escape(author)}</dc:creator>\n"
        f"<dc:language>{lang}</dc:language>\n"
        f'<meta property="dcterms:modified">{modified}</meta>\n'
        "</metadata>\n<manifest>\n"
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>\n'
        '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>\n'
        '<item id="css" href="style.css" media-type="text/css"/>\n'
        '<item id="title" href="title.xhtml" media-type="application/xhtml+xml"/>\n'
        f"{manifest}\n</manifest>\n"
        f'<spine toc="ncx">\n<itemref idref="title"/>\n{spine}\n</spine>\n</package>\n'
    )
    container = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
        "</rootfiles>\n</container>\n"
    )

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        # "mimetype" tiene que ser el primer archivo y sin comprimir.
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/toc.ncx", ncx, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/style.css", EPUB_CSS, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/title.xhtml", title_page, compress_type=zipfile.ZIP_DEFLATED)
        for name, _heading, content in chapters:
            zf.writestr(f"OEBPS/{name}", content, compress_type=zipfile.ZIP_DEFLATED)
    return buffer.getvalue()


def export_epub(project: Project, path: str | Path) -> None:
    atomic_write_bytes(path, build_epub_bytes(project))
